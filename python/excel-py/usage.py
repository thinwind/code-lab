from openpyxl import Workbook

wb = Workbook()
ws = wb.active

treeData = [
    ["Type", "Leaf Color", "Height"],
    ["Maple", "Red", 549],
    ["Oak", "Green", 783],
    ["Pine", "Green", 1204],
]
for row in treeData:
    ws.append(row)

from openpyxl.styles import Font

ft = Font(bold=True)
for row in ws["A1:C1"]:
    for cell in row:
        cell.font = ft

from openpyxl.chart import BarChart, Reference, Series

chart1 = BarChart()
chart1.type = "col"
chart1.title = "Tree Height"
chart1.y_axis.title = "Height(cm)"
chart1.x_axis.title = "Tree Type"
chart1.legend = None

data = Reference(ws, min_col=3, min_row=2, max_row=4, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)

chart1.add_data(data)
chart1.set_categories(categories)

ws.add_chart(chart1, "E1")
wb.save("TreeData.xlsx")
